#!/usr/bin/env python3
"""Build the teaching deck from the canonical Markdown and concise projections."""

from pathlib import Path
from copy import deepcopy
from datetime import datetime,timezone
import json,re,uuid,argparse,os,subprocess
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN,MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'LSSJ_skripta.md'
CONTENT=ROOT/'predavanja'
FONT='Arial'
FONT_FILE=None
BOLD_FILE=None
C={'paper':'F8F7F2','ink':'182D3D','muted':'546674','teal':'00695F','mint':'E7F2ED','amber':'9A570D','sand':'FAEDD8','line':'D9E1DD','white':'FFFFFF','pale':'EDF1F2','blue':'24557C'}
W,H=13.333333,7.5
WARN=[];BOXES=[]


def configure_fonts(regular=None,bold=None,family=None):
 """Locate installed Arial-compatible fonts; never require a workspace path."""
 global FONT_FILE,BOLD_FILE,FONT
 if bool(regular)!=bool(bold):
  raise ValueError('Navedite obe možnosti: --font-regular in --font-bold.')
 if regular:
  pair=(regular,bold)
 else:
  candidates=[]
  # Fontconfig covers Linux and installations on other operating systems.
  for name in ('Nimbus Sans','Arial','Liberation Sans'):
   matched=[]
   for style in ('Regular','Bold'):
    try:
     result=subprocess.run(['fc-match','-f','%{family}\n%{file}',f'{name}:style={style}'],
                           capture_output=True,text=True,check=True,timeout=10)
     actual,path=result.stdout.strip().split('\n',1)
     if name.casefold() not in actual.casefold():break
     matched.append(Path(path))
    except (OSError,ValueError,subprocess.SubprocessError):break
   if len(matched)==2 and matched[0]!=matched[1]:candidates.append(tuple(matched))
  # Native Windows and macOS installations need not have Fontconfig.
  windows=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts'
  mac=Path('/System/Library/Fonts/Supplemental')
  candidates.extend([(windows/'arial.ttf',windows/'arialbd.ttf'),
                     (mac/'Arial.ttf',mac/'Arial Bold.ttf')])
  pair=next((p for p in candidates if all(f.is_file() for f in p)),None)
  if pair is None:
   raise RuntimeError('Pisava za merjenje besedila ni najdena. Namestite Nimbus Sans '
                      '(fonts-urw-base35) ali podajte --font-regular in --font-bold. '
                      'Navodila: docs/UREJANJE.md.')
 for path in pair:
  if not Path(path).is_file():raise FileNotFoundError(path)
  ImageFont.truetype(str(path),40)
 FONT_FILE,BOLD_FILE=map(str,pair)
 FONT=family or 'Arial'
 _font_cache.clear()


def clean_source_block(text):
 """Remove Markdown separators, keeping all substantive source text."""
 return re.sub(r'^\s*---+\s*$', '',text,flags=re.M).strip()


def parse_items(text,section):
 """Read plain or bold numbered items without losing a bold item title."""
 pattern=re.compile(r'^(?:\*\*(\d+)\.\*\*[ \t]+|\*\*(\d+)\.[ \t]+|(\d+)\.[ \t]+)',re.M)
 markers=list(pattern.finditer(text));items=[]
 for i,m in enumerate(markers):
  end=markers[i+1].start() if i+1<len(markers) else len(text)
  original=clean_source_block(text[m.start():end])
  body=clean_source_block(text[m.end():end])
  if m[2]:body='**'+body
  number=int(m[1] or m[2] or m[3])
  if not body:raise ValueError(f'{section}: prazna naloga ali rešitev {number}.')
  items.append({'number':number,'text':body,'original':original,
                'hard':'Zahtevnejša naloga' in body})
 numbers=[item['number'] for item in items]
 if numbers!=list(range(1,len(items)+1)) or not items:
  raise ValueError(f'{section}: številčenje mora biti neprekinjeno od 1; najdeno {numbers}.')
 return items


def parse_source(master):
 """Derive chapters, sections, exercises and solutions from LSSJ_skripta.md."""
 headings=list(re.finditer(r'^## (\d+)\. (.+)$',master,re.M))
 if [int(m[1]) for m in headings]!=list(range(1,9)):
  raise ValueError('Skripta mora vsebovati poglavja 1–8 z naslovi oblike ## 1. Naslov.')
 appendix=re.search(r'^## Dodatek A\.',master,re.M)
 if appendix is None or '## Viri in literatura' not in master:
  raise ValueError('Manjkata dodatek A ali razdelek Viri in literatura.')
 source={}
 for i,h in enumerate(headings):
  chapter=h[1];end=headings[i+1].start() if i+1<len(headings) else appendix.start()
  body=clean_source_block(master[h.start():end])
  marks=list(re.finditer(r'^### (\d+\.\d+) (.+)$',body,re.M))
  if not marks:raise ValueError(f'Poglavje {chapter} nima oštevilčenih razdelkov.')
  sections={}
  for j,m in enumerate(marks):
   key=m[1]
   if key.split('.')[0]!=chapter or key in sections:
    raise ValueError(f'Neustrezna ali podvojena oznaka razdelka: {key}.')
   stop=marks[j+1].start() if j+1<len(marks) else len(body)
   sections[key]={'title':m[2],'text':clean_source_block(body[m.end():stop])}
  exercise_sections=[k for k,v in sections.items()
                     if v['title'].casefold() in ('vaje','naloge','sklepne vaje')]
  solution_sections=[k for k,v in sections.items()
                     if v['title'].casefold().startswith(('rešitve','predlagane rešitve'))]
  if len(exercise_sections)!=1 or len(solution_sections)!=1:
   raise ValueError(f'Poglavje {chapter} potrebuje po en razdelek vaj in rešitev.')
  es,ss=exercise_sections[0],solution_sections[0]
  exercises=parse_items(sections[es]['text'],es)
  solutions=parse_items(sections[ss]['text'],ss)
  if [x['number'] for x in exercises]!=[x['number'] for x in solutions]:
   raise ValueError(f'Poglavje {chapter}: vaje in rešitve se ne ujemajo po številkah.')
  source[chapter]={'title':h[2],'intro':clean_source_block(body[:marks[0].start()]),
                   'sections':sections,'exercise_section':es,'solution_section':ss,
                   'exercises':exercises,'solutions':solutions}
 return source,clean_source_block(master[:headings[0].start()])


def plain(t):
 t=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',str(t))
 t=re.sub(r'[*`#]','',t)
 return t.replace('  \n','\n').strip()

def notes_plain(t):
 t=re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'\1 — \2',str(t))
 t=re.sub(r'\[([^\]]+)\]\(#[^)]+\)',r'\1',t)
 return re.sub(r'[*`#]','',t).strip()

def inline(t):
 pat=re.compile(r'(\*\*[^*]+?\*\*|\*[^*\n]+?\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')
 result=[];pos=0
 for m in pat.finditer(t):
  if m.start()>pos:result.append((t[pos:m.start()],False,False,None))
  x=m[0]
  if x.startswith('**'):result.append((x[2:-2],True,False,None))
  elif x.startswith('*'):result.append((x[1:-1],False,True,None))
  elif x.startswith('`'):result.append((x[1:-1],False,False,None))
  else:
   q=re.match(r'\[([^\]]+)\]\(([^)]+)\)',x)
   result.append((q[1],False,False,q[2]))
  pos=m.end()
 if pos<len(t):result.append((t[pos:],False,False,None))
 return result or [('',False,False,None)]

_font_cache={}
def font(size,bold=False):
 if FONT_FILE is None:configure_fonts()
 key=(round(size*4),bold)
 if key not in _font_cache:_font_cache[key]=ImageFont.truetype(BOLD_FILE if bold else FONT_FILE,key[0])
 return _font_cache[key]

def lines_for(t,width,size,bold=False):
 f=font(size,bold); maxw=width*72*4
 result=0
 for para in plain(t).split('\n'):
  if not para:result+=.45;continue
  line='';n=1
  for word in para.split():
   candidate=(line+' '+word).strip()
   if f.getlength(candidate)>maxw and line:n+=1;line=word
   else:line=candidate
  result+=n
 return result

def required_height(t,w,size,bold=False,leading=1.15):
 return lines_for(t,w,size,bold)*size*leading/72+.06

def rgb(k):return RGBColor.from_string(C.get(k,k))

def rect(slide,x,y,w,h,fill,line=None,radius=False):
 # The comparison columns remain; their decorative card backgrounds do not.
 if radius and fill=='white':return None
 sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 if line and not radius:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.6)
 else:sh.line.fill.background()
 sh._element.spPr.append(OxmlElement('a:effectLst'))
 return sh

def tb(slide,text,x,y,w,h,size=24,minsize=None,bold=False,color='ink',italic=False,align=None,valign=MSO_ANCHOR.TOP,margin=.0,leading=1.12,tag=''):
 size=float(size);minsize=float(minsize or size)
 avail_w=w-2*margin;avail_h=h-2*margin
 while size>minsize and required_height(text,avail_w,size,bold,leading)>avail_h:size-=.5
 need=required_height(text,avail_w,size,bold,leading)
 if need>avail_h+.025:WARN.append({'slide':slide._idx,'tag':tag,'chars':len(plain(text)),'height':h,'estimated_height':round(need,3),'font':size,'text':plain(text)[:110]})
 sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
 tf.margin_left=tf.margin_right=Inches(margin);tf.margin_top=tf.margin_bottom=Inches(margin);tf.vertical_anchor=valign
 for i,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
  p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.10
  if align is not None:p.alignment=align
  for chunk,b,it,url in inline(line):
   r=p.add_run();r.text=chunk
   r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold or b;r.font.italic=italic or it;r.font.color.rgb=rgb(color)
   if url and url.startswith('http'):r.hyperlink.address=url
  pPr=p._p.get_or_add_pPr();pPr.set('rtl','0')
 BOXES.append({'slide':slide._idx,'tag':tag,'x':x,'y':y,'w':w,'h':h,'font':size,'text':plain(text)})
 return sh

def small_caps(slide,text,x,y,w,color='teal'):
 return tb(slide,text.upper(),x,y,w,.28,11.5,bold=True,color=color,tag='eyebrow')

def refs_display(refs):
 refs=list(dict.fromkeys(refs or []))
 if not refs:return ''
 text=' · '.join(refs)
 if len(text)>182:return text[:179].rsplit(' ',1)[0]+' …'
 return text

def frame(slide,spec,index,total,targets):
 dark=spec['kind'] in ('cover','chapter','appendix_open')
 slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb('paper')
 if dark:return
 tb(slide,spec['title'],.66,.70,12.0,1.10,33,29,bold=True,tag='title')
 tb(slide,str(index),12.03,7.005,.65,.22,10,color='muted',align=PP_ALIGN.RIGHT,tag='slide-number')

def takeaway(slide,text):
 if not text:return
 tb(slide,text,.68,6.24,11.98,.47,18.5,17.5,bold=True,color='teal',tag='takeaway')

def concept(slide,s):
 lead=s.get('lead','')
 if lead:tb(slide,lead,.68,1.87,11.98,.82,24,22,color='muted',tag='lead')
 blocks=s.get('blocks',[]);n=len(blocks);gap=.24;w=(11.98-gap*(n-1))/n
 y=2.74 if lead else 2.05;h=6.0-y
 for i,b in enumerate(blocks):
  x=.68+i*(w+gap);rect(slide,x,y,w,h,'white','line',True)
  tb(slide,b['label'],x+.22,y+.23,w-.44,.62,17,15.5,bold=True,color='teal',tag='concept-label')
  tb(slide,b['text'],x+.22,y+.88,w-.44,h-1.06,25,21,tag='concept-text')
 takeaway(slide,s.get('takeaway'))

def contrast(slide,s):
 for i,k in enumerate(['left','right']):
  item=s[k];x=.68+i*6.10;y=2.02;w=5.88
  rect(slide,x,y,w,4.00,'white','line',True)
  tb(slide,item['label'],x+.23,y+.22,w-.46,.58,16,14.5,bold=True,color='teal',tag='contrast-label')
  tb(slide,item['example'],x+.23,y+.99,w-.46,1.78,27,22,bold=False,italic=True,tag='contrast-example')
  tb(slide,item.get('text',''),x+.23,y+2.87,w-.46,.98,20.5,19.5,color='muted',tag='contrast-explanation')
 takeaway(slide,s.get('takeaway'))

def example(slide,s):
 rect(slide,.68,1.98,11.98,2.08,'pale',radius=True)
 tb(slide,s['example'],.98,2.21,11.36,1.68,28,23,italic=True,tag='worked-example')
 steps=s.get('steps',[]);n=len(steps);gap=.24;w=(11.98-gap*(n-1))/n
 for i,b in enumerate(steps):
  x=.68+i*(w+gap)
  tb(slide,b['label'],x,4.24,w,.37,15.5,14.5,bold=True,color='teal',tag='example-label')
  tb(slide,b['text'],x,4.72,w,1.40,21.5,20,tag='example-step')
 takeaway(slide,s.get('takeaway'))

def process(slide,s):
 steps=s['steps'];n=len(steps)
 if n==4:
  for i,b in enumerate(steps):
   col=i%2;row=i//2;x=.68+col*6.10;y=2.01+row*2.06;w=5.88;h=1.89
   rect(slide,x,y,w,h,'white','line',True)
   tb(slide,str(i+1),x+.20,y+.18,.57,.46,24,bold=True,color='teal',tag='process-number')
   tb(slide,b['label'],x+.99,y+.19,w-1.22,.46,18.5,17,bold=True,tag='process-label')
   tb(slide,b['text'],x+.23,y+.79,w-.46,.91,22,20.5,tag='process-text')
 else:
  gap=.17;h=(3.93-gap*(n-1))/n
  for i,b in enumerate(steps):
   y=2.05+i*(h+gap);rect(slide,.68,y,11.98,h,'white','line',True)
   if not b['label'].startswith('Korak '):tb(slide,str(i+1),.91,y+.15,.70,h-.22,28,bold=True,color='teal',tag='process-number')
   tb(slide,b['label'],1.90,y+.12,3.1,h-.20,20.5,18,bold=True,tag='process-label')
   tb(slide,b['text'],5.30,y+.08,7.05,h-.11,23,21,tag='process-text')
 takeaway(slide,s.get('takeaway'))


def table(slide,s):
 headers=s['headers'];rows=s['rows'];n=len(headers)
 widths=s.get('widths') or ([4.1,7.88] if n==2 else [3.50,4.10,4.38])
 size=22 if n==2 else 20.5
 heights=[]
 for row in rows:
  heights.append(max(.82,max(required_height(str(v),widths[j]-.34,size) for j,v in enumerate(row))+.20))
 total=.51+sum(heights)
 while total>4.2 and size>19.5:
  size-=.5
  heights=[max(.80,max(required_height(str(v),widths[j]-.34,size) for j,v in enumerate(row))+.18) for row in rows]
  total=.51+sum(heights)
 if total>4.27:WARN.append({'slide':slide._idx,'tag':'table-total','height':round(total,2),'title':s['title']})
 tab=slide.shapes.add_table(len(rows)+1,n,Inches(.68),Inches(1.98),Inches(11.98),Inches(total)).table
 for j,w in enumerate(widths):tab.columns[j].width=Inches(w)
 for i,row in enumerate([headers]+rows):
  tab.rows[i].height=Inches(.51 if i==0 else heights[i-1])
  for j,value in enumerate(row):
   cell=tab.cell(i,j);cell.margin_left=cell.margin_right=Inches(.15);cell.margin_top=cell.margin_bottom=Inches(.08)
   cell.vertical_anchor=MSO_ANCHOR.MIDDLE;cell.fill.solid();cell.fill.fore_color.rgb=rgb('teal' if i==0 else ('white' if i%2 else 'pale'))
   tf=cell.text_frame;tf.clear();tf.word_wrap=True
   for k,line in enumerate(str(value).split('\n')):
    p=tf.paragraphs[0] if k==0 else tf.add_paragraph();p.space_after=Pt(0);p.line_spacing=1.08
    for text,b,it,url in inline(line):
     r=p.add_run();r.text=text;r.font.name=FONT;r.font.size=Pt(15.5 if i==0 else size);r.font.bold=(i==0 or b);r.font.italic=it;r.font.color.rgb=rgb('white' if i==0 else 'ink')
     if url and url.startswith('http'):r.hyperlink.address=url
 takeaway(slide,s.get('takeaway'))

def practice(slide,s):
 items=s['items'];n=len(items);answer=s['kind']=='answer';col='teal' if answer else 'amber'
 if n==2:
  for i,item in enumerate(items):
   x=.68+i*6.10;w=5.88
   rect(slide,x,2.0,w,4.19,'white','line',True)
   tb(slide,str(item['number'])+'.',x+.22,2.20,w-.44,.35,18,bold=True,color=col,tag='practice-number')
   tb(slide,item['text'],x+.22,2.66,w-.44,3.33,23,20,tag='answer-text' if answer else 'exercise-text')
 else:
  item=items[0];rect(slide,.68,2.0,11.98,4.20,'white','line',True)
  tb(slide,item['text'],1.00,2.16,11.34,3.88,25,22,tag='answer-text' if answer else 'exercise-text')


def opening(slide,s,index,total):
 if s['kind']=='cover':
  tb(slide,'Leksika in slovnica\nslovenskega jezika',.74,1.52,11.80,1.70,47,bold=True,tag='cover-title')
  tb(slide,'Od poimenovanja in zgradbe\ndo utemeljenega jezikovnega nasveta',.78,3.63,11.72,1.2,30,color='muted',tag='cover-subtitle')
  tb(slide,'Damjan Popič',.78,6.45,11.6,.40,18,color='muted',tag='cover-author')
 else:
  title=(str(s['chapter'])+'. ' if s['kind']=='chapter' else '')+s['title']
  tb(slide,title,.76,1.35,11.6,1.45,39,34,bold=True,tag='chapter-title')
  tb(slide,s.get('question',''),.80,3.20,11.5,1.00,26,24,color='muted',tag='chapter-question')
  outs=s.get('outcomes',[])
  for i,line in enumerate(outs):
   tb(slide,'• '+line,.80,4.65+i*.61,11.5,.52,22,20,tag='chapter-outcome')
 if s['kind']!='cover':tb(slide,str(index),12.03,7.005,.65,.22,10,color='muted',align=PP_ALIGN.RIGHT,tag='slide-number')

def recap(slide,s):
 for i,line in enumerate(s['points']):
  y=2.02+i*1.02
  tb(slide,str(i+1),.72,y,.68,.63,29,bold=True,color='teal',tag='recap-number')
  tb(slide,line,1.65,y+.04,10.80,.76,25,23,tag='recap-text')
 rect(slide,.68,5.31,11.98,.77,'pale',radius=True)
 tb(slide,s.get('bridge',''),.92,5.47,11.49,.53,19.5,18,tag='recap-bridge')
 if s.get('reading'):tb(slide,'Branje: '+s['reading'],.78,6.29,11.60,.47,13.5,12.5,color='muted',tag='reading-pointers')

def toc(slide,s,targets,ranges):
 for i,n in enumerate(range(1,9)):
  col=0 if i<4 else 1;row=i%4;x=.68+col*6.10;y=1.99+row*1.04;w=5.88
  tb(slide,str(n)+'.',x+.05,y+.18,.60,.48,25,bold=True,color='teal',tag='overview-number')
  tb(slide,s['entries'][n-1]['short_title'],x+.80,y+.16,4.90,.65,24,22,tag='overview-title')


def glossary(slide,s):
 items=s['items'];n=len(items);gap=.15;h=(4.18-gap*(n-1))/n
 for i,item in enumerate(items):
  y=1.98+i*(h+gap);rect(slide,.68,y,11.98,h,'white','line',True)
  tb(slide,item['term'],.88,y+.14,3.5,h-.24,19.5,17.5,bold=True,color='teal',tag='glossary-term')
  tb(slide,item['definition'],4.60,y+.14,7.82,h-.24,20.5,19.5,tag='glossary-definition')


def bibliography(slide,s):
 items=s['items'];n=len(items);gap=.20;h=(4.15-gap*(n-1))/n
 for i,item in enumerate(items):
  y=2.01+i*(h+gap)
  text=item['title']
  sh=tb(slide,text,.78,y,11.83,h,21.5,19,tag='bibliography-entry')
  if item.get('url'):
   sh.click_action.hyperlink.address=item['url']


def make_notes(slide,s,source,covered):
 parts=[f"LSSJ · {s['title']}"]
 if s.get('section'):parts.append('Povezava s skripto: '+s['section'])
 if s.get('notes'):parts.append('ZA IZVEDBO\n'+s['notes'])
 if s.get('refs'):parts.append('VIRI\n'+'\n'.join(s['refs']))
 ch=str(s.get('chapter',0));c=source.get(ch)
 if c:
  for section in s.get('source_sections',[]):
   if section not in covered and section in c['sections']:
    v=c['sections'][section];parts.append(f"RAZLAGA IZ SKRIPTE · {section} {v['title']}\n"+v['text']);covered.add(section)
  if s['kind']=='chapter':parts.append('UVOD V ENOTO\n'+c['intro'])
 if s.get('full_source'):parts.append('BESEDILO IZ SKRIPTE\n'+s['full_source'])
 if s['kind'] in ('exercise','answer'):
  key='NALOGE' if s['kind']=='exercise' else 'REŠITVE Z RAZLAGO'
  parts.append(key+'\n\n'+'\n\n'.join(item['original'] for item in s['items']))
 text=notes_plain('\n\n'.join(parts))
 slide.notes_slide.notes_text_frame.text=text
 return len(text.split())

def markdown_table(text):
 lines=[l.strip() for l in text.splitlines() if l.strip().startswith('|')]
 if not lines:return []
 out=[]
 for line in lines[2:]:
  cells=[x.strip() for x in line.strip('|').split('|')]
  out.append(cells)
 return out

def practice_specs(chapter,c):
 ex=c['exercises'];sol={x['number']:x for x in c['solutions']};groups=[];i=0
 while i<len(ex):
  group=[ex[i]]
  if i+1<len(ex) and not ex[i]['hard'] and not ex[i+1]['hard']:
   a,b=sol[ex[i]['number']],sol[ex[i+1]['number']]
   if len(plain(a['text']).split())+len(plain(b['text']).split())<=130 and max(required_height(a['text'],5.44,20),required_height(b['text'],5.44,20))<=3.28:
    group.append(ex[i+1])
  groups.append(group);i+=len(group)
 specs=[]
 for group in groups:
  nums=[x['number'] for x in group];label=' in '.join(map(str,nums));hard=any(x['hard'] for x in group)
  key=f'c{chapter}-e'+('-'.join(map(str,nums)))
  prompt={'id':key,'kind':'exercise','chapter':chapter,'title':('Zahtevnejša vaja ' if hard else ('Vaji ' if len(nums)==2 else 'Vaja '))+label,'section':c['exercise_section'],'items':group,'refs':[f"Skripta 3.0, {c['exercise_section']}; "+('nalogi ' if len(nums)==2 else 'naloga ')+label],'notes':'Najprej omogočite samostojen poskus. Zahtevajte rešitev in odločilni razlog. Pri več mogočih analizah naj študent izrecno navede predpostavko. Razlaga je na naslednji prosojnici.','pair_id':key+'-answer','hard':hard}
  answer={'id':key+'-answer','kind':'answer','chapter':chapter,'title':('Rešitvi ' if len(nums)==2 else 'Rešitev ')+label+' z razlago','section':c['solution_section'],'items':[sol[n] for n in nums],'refs':[f"Skripta 3.0, {c['solution_section']}; "+('rešitvi ' if len(nums)==2 else 'rešitev ')+label],'notes':'Primerjajte razlago s postopkom študentov. Naj označijo, kaj so dokazali in katera predpostavka je odločilna. Alternative sprejmite v mejah, navedenih v skripti.','pair_id':key,'hard':hard}
  specs.extend([prompt,answer])
 return specs

def make_appendices(master):
 text=master.split('## Dodatek A.',1)[1]
 text='## Dodatek A.'+text.split('## Viri in literatura',1)[0]
 specs=[]
 parts={}
 headings=list(re.finditer(r'^## Dodatek ([A-D])\. (.+)$',text,re.M))
 for i,m in enumerate(headings):parts[m[1]]={'title':m[2],'text':text[m.end():headings[i+1].start() if i+1<len(headings) else len(text)].strip()}
 specs.append({'id':'appendix-start','kind':'appendix_open','group':'09 · Dodatki','title':'Pripomočki za ponavljanje','letter':'A–D','question':'Pojem povežemo s primerom, postopkom in virom.','outcomes':['Pojmovni pregled','Postopki in izpitni problemi','Usmerjeno branje'], 'notes':'Dodatki so namenjeni ponavljanju in sprotnemu vračanju k pojmom. Lahko jih uporabite ob posamezni enoti ali na koncu. Ob kratki definiciji vedno zahtevajte konkreten zgled.'})
 terms=markdown_table(parts['A']['text']);i=0;page=0
 while i<len(terms):
  batch=terms[i:i+4]
  while len(batch)>2:
   h=(4.18-.15*(len(batch)-1))/len(batch)-.24
   if all(max(required_height(r[0],3.5,17.5,True),required_height(r[1],7.82,19.5))<=h for r in batch):break
   batch=batch[:-1]
  page+=1
  specs.append({'id':f'glossary-{page}','kind':'glossary','group':'09 · Dodatek A — Pojmi','title':f'Pojmovni pregled · {page:02d}','items':[{'term':r[0],'definition':r[1]} for r in batch],'refs':['Skripta 3.0, dodatek A'],'notes':'Definicije so orientacija za ponavljanje. Izberite pojem, povejte zgled in pokažite mejo njegove uporabe. Pri teoretično odvisnih pojmih navedite model. Poglavja za nadaljnje branje so v izvirnih vrsticah spodaj.','full_source':'\n'.join(' | '.join(r) for r in batch)})
  i+=len(batch)
 # Repetition procedures, kept in their original order.
 bheads=list(re.finditer(r'^### (B\.\d+) (.+)$',parts['B']['text'],re.M))
 for k,m in enumerate(bheads):
  body=parts['B']['text'][m.end():bheads[k+1].start() if k+1<len(bheads) else len(parts['B']['text'])].strip()
  items=re.findall(r'^\d+\. (.+)$',body,re.M)
  starts=[0,4] if len(items)==7 else list(range(0,len(items),3))
  for start in starts:
   count=4 if len(items)==7 and start==0 else 3
   part=items[start:start+count]
   specs.append({'id':f'procedure-{m[1]}-{start}','kind':'process','group':'10 · Dodatek B — Postopki','title':m[2]+(' · nadaljevanje' if start else ''),'steps':[{'label':f'Korak {start+j+1}','text':line} for j,line in enumerate(part)],'takeaway':'Vsak korak uporabimo na konkretnem zgledu.','refs':[f'Skripta 3.0, dodatek {m[1]}'],'notes':'To je seznam vprašanj za ponavljanje, ne nadomestilo za razčlenitev. Pri novem primeru izberite odločilni korak in pojasnite, kateri podatek potrebujete.','full_source':body if start==0 else ''})
 # Exam problem families: compact mapping, full original mapping in notes.
 mapping=[['Vir in jezikovni nasvet','Isti pomen, isto vprašanje, jasen status vira.','1 in 8'],['Lastno, občno, opis','Identifikacija, poimenovalna funkcija in začetnica.','2 in 5'],['Tvorjenka in zapis','SPo, BPo, obrazilo; nato pravopisni okvir.','3'],['Oblika v sobesedilu','Sklon, ujemanje, navezava, vid in funkcija.','4'],['Tuja imena','Izgovor, osnova, pregibanje in tvorjenje.','5'],['Vejica ob povezovalcu','Hierarhija sestavin in obe meji odvisnika.','6'],['Posebne vključitve','Polstavek, pristavek, vrivek, nagovor in navedek.','7'],['Celovito urejanje','Pomenska zvestoba in dokaz za posamezni poseg.','8']]
 for start in [0,4]:
  specs.append({'id':f'exam-map-{start}','kind':'table','group':'11 · Dodatek C — Izpitni problemi','title':'Izpitni problem povežemo s postopkom'+(' · 2' if start else ' · 1'),'headers':['Družina nalog','Odločilno vprašanje','Enota'],'rows':mapping[start:start+4],'widths':[3.9,6.6,1.48],'takeaway':'Pri novem primeru najprej prepoznamo skupni problem, nato razlikovalni podatek.','refs':['Skripta 3.0, dodatek C'],'notes':'Pregled arhiva usmerja izbor družin problemov. Ne gre za uradni točkovnik ali napoved konkretnega izpita. Prenos pokažemo s spremenjenim primerom, ne s ponavljanjem istega odgovora.','full_source':parts['C']['text'] if start==0 else ''})
 readings=markdown_table(parts['D']['text'])
 for i,row in enumerate(readings):
  if len(row)<3:continue
  specs.append({'id':f'reading-{i+1}','kind':'concept','group':'12 · Dodatek D — Branje','title':'Usmerjeno branje · '+row[0],'lead':'Od avtorjevega merila do lastne razčlenitve.','blocks':[{'label':'Odlomki','text':row[1]},{'label':'Vprašanje po branju','text':row[2]}],'takeaway':'Sklici navajajo natisnjene strani izdaj, opredeljenih v bibliografiji.','refs':['Skripta 3.0, dodatek D'],'notes':'Preberite navedeni odlomek, izpišite merilo in preverite en avtorjev ter en lasten zgled. Pri dopolnjenih izdajah vključite opombe. Razpon je načrt problemskega branja, ne obveza, da vsakokrat preberete vse strani v enem zamahu.','full_source':' | '.join(row)})
 specs.append({'id':'reading-record','kind':'process','group':'12 · Dodatek D — Branje','title':'Bralni zapis: od navedbe do dokaza','steps':[{'label':'Trditev in stran','text':'S svojimi besedami zapišem avtorjevo merilo in točno stran.'},{'label':'Zgled in prenos','text':'Razčlenim avtorjev zgled in uporabim merilo na novem.'},{'label':'Meja in posledica','text':'Povem, do kod sklep velja in kaj spremeni pri jezikovni odločitvi.'}],'takeaway':'Primerjajmo najmočnejši in najšibkejši dokaz za svojo rešitev.','refs':['Skripta 3.0, dodatek D.1'],'notes':'V paru preverita, katera trditev v zapisu je prepis kategorije in katera je že razčlenitev. Izberita protiprimer, ki pokaže mejo uporabe merila.','full_source':parts['D']['text'][parts['D']['text'].find('### D.1'):]})
 bib=master.split('## Viri in literatura',1)[1]
 specs.append({'id':'books','kind':'bibliography','group':'13 · Viri in literatura','title':'Tri temeljna dela','items':[
  {'title':'Jože Toporišič (2004)\nSlovenska slovnica. 4., prenovljena in razširjena izdaja, 2. natis. Maribor: Obzorja.'},
  {'title':'Ada Vidovič Muha (2011)\nSlovensko skladenjsko besedotvorje. 2., razširjena in dopolnjena izdaja. Ljubljana: Znanstvena založba FF UL.'},
  {'title':'Ada Vidovič Muha (2013)\nSlovensko leksikalno pomenoslovje. 2., dopolnjena izdaja. Ljubljana: Znanstvena založba FF UL.'}], 'refs':['Navedbe sledijo navedenim izdajam in natisu.'],'notes':'Toporišič 2004 je drugi natis četrte izdaje iz leta 2000. Knjižni sklici uporabljajo natisnjene strani. Celotna bibliografija skripte je spodaj.','full_source':bib})
 specs.append({'id':'norm-sources','kind':'bibliography','group':'13 · Viri in literatura','title':'Pravopisni viri: navedemo tudi status','items':[
  {'title':'Slovenski pravopis 2001\nPravila: sklic na člen (§). Slovarske sestavke navajamo posebej.','url':'https://www.fran.si/134/slovenski-pravopis/datoteke/Pravopis_Pravila.pdf'},
  {'title':'ePravopis 2025 (objavljen 2026)\nElektronski slovarski zvezek; predlog. Sklic na iztočnico in stran.','url':'https://doi.org/10.3986/9789610511366'},
  {'title':'Pravopis 8.0 — predlog, 22. junij 2026\nSklic na poglavje, člen in PDF-stran.','url':'https://slovenski-pravopis.si/wp-content/uploads/2026/06/2026-06-22_SP8.pdf'}], 'refs':['Skripta 3.0: normativni okvir preverjen 6. oktobra 2026.'],'notes':'SP2001, predlagane slovarske rešitve in predlog pravil imajo v skripti jasno razmejen status. Predloga ne predstavimo kot že sprejetega pravila; staro rešitev ocenimo tudi v njenem takratnem okviru.'})
 web=[]
 segment=bib.split('### Strokovne obravnave in konkretni slovarski sestavki',1)[-1]
 for line in segment.splitlines():
  if not line.startswith('- '):continue
  m=re.search(r'\[([^\]]+)\]\((https?://[^)]+)\)',line)
  if not m:continue
  label=plain(line[2:]);web.append({'title':label,'url':m[2]})
 for start in range(0,len(web),4):
  specs.append({'id':f'web-refs-{start}','kind':'bibliography','group':'13 · Viri in literatura','title':f'Spletni viri · {start//4+1:02d}','items':web[start:start+4],'refs':['Povezave so klikljive; polne navedbe so tudi v opombah.'],'notes':'Odprite naslov za konkretni vir. Pri odgovoru preverite, ali vir obravnava isti pomen, zgradbo in normativni okvir.','full_source':'\n'.join(x['title']+'\n'+x['url'] for x in web[start:start+4])})
 return specs,len(terms)


def expand_table_specs(specs):
 out=[]
 for s in specs:
  if s['kind']!='table':out.append(s);continue
  n=len(s['headers']);widths=s.get('widths') or ([4.1,7.88] if n==2 else [3.5,4.1,4.38]);rows=s['rows']
  h=.51+sum(max(.80,max(required_height(str(v),widths[j]-.34,19.5) for j,v in enumerate(row))+.18) for row in rows)
  if h<=4.25:out.append(s);continue
  batches=[];current=[]
  for row in rows:
   test=current+[row];height=.51+sum(max(.80,max(required_height(str(v),widths[j]-.34,19.5) for j,v in enumerate(r))+.18) for r in test)
   if current and height>4.2:batches.append(current);current=[row]
   else:current=test
  if current:batches.append(current)
  for i,b in enumerate(batches):
   cp=deepcopy(s);cp['id']=s['id']+f'-part{i+1}';cp['rows']=b
   if i:cp['title']+=' · nadaljevanje'
   out.append(cp)
 return out


def split_answer_specs(specs):
 out=[]
 for s in specs:
  if s['kind']!='answer' or len(s['items'])!=1 or required_height(s['items'][0]['text'],11.34,22)<=3.87:
   out.append(s);continue
  original=s['items'][0]['text'];candidates=[]
  for m in re.finditer(r'(?<=[.!?])\s+(?=[A-ZČŠŽ(])',original):
   if original[:m.start()].count('*')%2==0:candidates.append(m.start())
  if not candidates:
   for m in re.finditer(r'\s+',original):
    if original[:m.start()].count('*')%2==0:candidates.append(m.start())
  middle=min(candidates,key=lambda pos:abs(pos-len(original)/2))
  chunks=[original[:middle].strip(),original[middle:].strip()]
  assert ' '.join(chunks).split()==original.split()
  for i,chunk in enumerate(chunks):
   cp=deepcopy(s);cp['id']=s['id'] if i==0 else s['id']+'-part2'
   cp['title']=s['title']+f' · {i+1}/2';cp['items'][0]['text']=chunk
   cp['notes']=s['notes']+f' Daljša razlaga je razdeljena na dva dela; to je del {i+1}.'
   out.append(cp)
 return out


def collect_specs(sample=False):
 master=SOURCE.read_text(encoding='utf-8')
 source,intro=parse_source(master)
 contents=[]
 for n in range(1,9):
  p=CONTENT/f'chapter_{n}.json'
  if not p.exists():
   if sample:continue
   raise FileNotFoundError(p)
  data=json.loads(p.read_text(encoding='utf-8'))
  if data.get('chapter')!=n:raise ValueError(f'{p.name}: napačna številka poglavja.')
  data['title']=source[str(n)]['title']
  for slide in data['slides']:
   unknown=set(slide.get('source_sections',[]))-set(source[str(n)]['sections'])
   if unknown:raise ValueError(f"{slide['id']}: neobstoječi razdelki {sorted(unknown)}.")
  contents.append(data)
 specs=[{'id':'cover','kind':'cover','group':'00 · Uvod','title':'Leksika in slovnica slovenskega jezika','notes':'Predavanja po teoretično poglobljeni skripti 3.0. Vsebine so razporejene v osem poglavij. Vajam sledijo ločene rešitve. Razlage iz skripte in natančni sklici so vključeni v opombe ustreznih prosojnic.','full_source':intro},
 {'id':'toc','kind':'toc','group':'00 · Uvod','title':'Vsebina','entries':contents,'notes':'Vsebine izvajamo po zaporedju. Vsako poglavje vsebuje razlage in primere, vaje z rešitvami ter povzetek.'},
 {'id':'guide','kind':'concept','group':'00 · Uvod','title':'Kako bomo delali?','lead':'Jezikovno odločitev oblikujemo, pokažemo in preverimo.','blocks':[{'label':'Razčlenimo','text':'Prepoznamo pomen, sestavine in odločilno razmerje.'},{'label':'Utemeljimo','text':'Rešitev podpremo s postopkom in ustreznim virom.'},{'label':'Prenesemo','text':'Spremenimo en podatek in preverimo, ali prejšnja razlaga še velja.'}],'takeaway':'Pri vajah primerjamo postopek, ne samo končnega zapisa.','notes':'Pojasnite način dela: najprej vprašanje, nato poskus in razlaga. Zahtevnejše naloge so del gradiva in so označene. Po razkritju rešitve naj študent pove, kateri dokaz je bil odločilen. Ocenjevanje konkretnega izpita sledi predavateljevim navodilom.'}]
 for data in contents:
  n=data['chapter'];c=source[str(n)];group=f'{n:02d} · {data["short_title"]}'
  specs.append({'id':f'chapter-{n}','kind':'chapter','chapter':n,'group':group,'title':data['title'],'question':data['question'],'outcomes':data['outcomes'],'notes':'Osrednje vprašanje te enote: '+data['question']+'\nCilji: '+'; '.join(data['outcomes'])})
  for s in data['slides']:
   cp=deepcopy(s);cp['chapter']=n;cp['group']=group;specs.append(cp)
  p=practice_specs(n,c)
  for s in p:s['group']=group
  specs.extend(p)
  after_sections=[k for k in c['sections'] if int(k.split('.')[1])>int(c['solution_section'].split('.')[1])]
  wrap=data['wrap'];specs.append({'id':f'chapter-{n}-recap','kind':'recap','chapter':n,'group':group,'title':'Kaj prenesemo v naslednji primer?','points':wrap['points'],'bridge':wrap['bridge'],'reading':wrap['reading'],'source_sections':after_sections,'notes':'Pred prehodom na novo enoto naj študenti za vsako ugotovitev navedejo svoj primer. Nato preverimo, kaj bi moralo biti drugače, da bi se spremenila rešitev.','refs':[f'Skripta3.0, poglavje{n}']})
 apps,terms=make_appendices(master);specs.extend(apps)
 specs=expand_table_specs(split_answer_specs(specs))
 # Plain visible headings; detailed section identifiers remain in the notes.
 chapter_titles={d['chapter']:d.get('display_title',d['short_title']) for d in contents}
 for s in specs:
  if s['kind']=='chapter':s['title']=chapter_titles[s['chapter']]
  elif s['kind']=='recap':s['title']='Povzetek'
  elif s['kind']=='answer':
   s['title']=s['title'].replace(' z razlago','').replace(' · 1/2','').replace(' · 2/2',' (nadaljevanje)')
  elif s['kind']=='glossary':s['title']='Pojmovni pregled'
  elif s['id'].startswith('web-refs-'):s['title']='Spletni viri'
  elif s['id'].startswith('reading-') and s['id']!='reading-record':
   s['title']=re.sub(r'^Usmerjeno branje · \d+\. ', 'Branje: ',s['title'])
  elif s['id'].startswith('exam-map-'):s['title']='Izpitni problemi'
  elif s['id']=='norm-sources':s['title']='Pravopisni viri'
  s['title']=s['title'].replace(' · nadaljevanje',' (nadaljevanje)')
  if s.get('group'):s['group']=re.sub(r'^0*(\d+) · ',r'\1. ',s['group'])
  if s.get('chapter'):s['group']=str(s['chapter'])+'. '+chapter_titles[s['chapter']]
 if sample:
  keep=[];seen=set()
  for s in specs:
   if s['kind'] in ['cover','chapter','concept','contrast','example','table','process','exercise','answer','recap','glossary','bibliography'] and s['kind'] not in seen:
    keep.append(s);seen.add(s['kind'])
  specs=keep
 return specs,source,terms


def add_sections(prs,specs):
 nsP='http://schemas.openxmlformats.org/presentationml/2006/main'
 ns14='http://schemas.microsoft.com/office/powerpoint/2010/main'
 el=prs._element
 extlst=el.find('{'+nsP+'}extLst')
 if extlst is None:extlst=etree.SubElement(el,'{'+nsP+'}extLst')
 ext=etree.SubElement(extlst,'{'+nsP+'}ext',uri='{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
 lst=etree.SubElement(ext,'{'+ns14+'}sectionLst',nsmap={'p14':ns14})
 last=None;slst=None
 for s,slide in zip(specs,prs.slides):
  group=s.get('group','Gradivo')
  if group!=last:
   sec=etree.SubElement(lst,'{'+ns14+'}section',name=group,id='{'+str(uuid.uuid5(uuid.NAMESPACE_URL,'lssj-2026/'+group)).upper()+'}')
   slst=etree.SubElement(sec,'{'+ns14+'}sldIdLst');last=group
  etree.SubElement(slst,'{'+ns14+'}sldId',id=str(slide.slide_id))


def normalize_app_metadata(prs,specs):
 """Replace the empty template's counts and 4:3 label with real deck metadata."""
 ns='http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
 vt='http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
 part=next(p for p in prs.part.package.iter_parts() if str(p.partname)=='/docProps/app.xml')
 root=etree.fromstring(part.blob)
 def put(name,value):
  element=root.find('{'+ns+'}'+name)
  if element is None:element=etree.SubElement(root,'{'+ns+'}'+name)
  element.text=str(value)
 count=len(prs.slides)
 put('Application','python-pptx')
 put('PresentationFormat','Widescreen (16:9)')
 put('Slides',count)
 put('Notes',sum(slide.has_notes_slide for slide in prs.slides))
 put('HiddenSlides',0)
 put('TotalTime',0)
 words=paragraphs=0
 for slide in prs.slides:
  for shape in slide.shapes:
   frames=[]
   if shape.has_text_frame:frames.append(shape.text_frame)
   if shape.has_table:frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
   for frame in frames:
    words+=len(frame.text.split())
    paragraphs+=len(frame.paragraphs)
 put('Words',words)
 put('Paragraphs',paragraphs)
 # These two lists describe the theme and every actual slide title.
 for name in ('HeadingPairs','TitlesOfParts'):
  element=root.find('{'+ns+'}'+name)
  if element is not None:root.remove(element)
 pairs=etree.SubElement(root,'{'+ns+'}HeadingPairs')
 vector=etree.SubElement(pairs,'{'+vt+'}vector',size='4',baseType='variant')
 for tag,value in [('lpstr','Theme'),('i4','1'),('lpstr','Slide Titles'),('i4',str(count))]:
  variant=etree.SubElement(vector,'{'+vt+'}variant')
  etree.SubElement(variant,'{'+vt+'}'+tag).text=value
 titles=etree.SubElement(root,'{'+ns+'}TitlesOfParts')
 vector=etree.SubElement(titles,'{'+vt+'}vector',size=str(count+1),baseType='lpstr')
 for title in ['Office Theme']+[s['title'] for s in specs]:
  etree.SubElement(vector,'{'+vt+'}lpstr').text=title
 # AppVersion belongs to the original template's PowerPoint application.
 app_version=root.find('{'+ns+'}AppVersion')
 if app_version is not None:root.remove(app_version)
 part._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)


def main():
 ap=argparse.ArgumentParser(description='Obnovi predavanja iz skripte in projekcij poglavij.')
 ap.add_argument('--sample',action='store_true',help='Samo vzorec vrst prosojnic.')
 ap.add_argument('--output',type=Path,help='Izhodna datoteka PPTX; privzeto v korenu repozitorija.')
 ap.add_argument('--report',type=Path,help='Po želji shrani poročilo o vsebini in postavitvi v JSON.')
 ap.add_argument('--font-regular',type=Path,help='Datoteka običajne pisave za merjenje besedila.')
 ap.add_argument('--font-bold',type=Path,help='Datoteka krepke različice iste pisave.')
 ap.add_argument('--font-family',help='Ime pisave, zapisano v PPTX; privzeto samodejno.')
 args=ap.parse_args()
 configure_fonts(args.font_regular,args.font_bold,args.font_family)
 specs,source,termcount=collect_specs(args.sample)
 prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
 prs.core_properties.title='Leksika in slovnica slovenskega jezika'
 prs.core_properties.subject='Celovit predavateljski komplet po skripti3.0'
 prs.core_properties.author='Damjan Popič';prs.core_properties.keywords='LSSJ, slovnica, leksika, Toporišič, Vidovič Muha, pravopis'
 prs.core_properties.last_modified_by='Damjan Popič'
 prs.core_properties.created=datetime.now(timezone.utc)
 prs.core_properties.modified=prs.core_properties.created
 prs.core_properties.comments='Po teoretično poglobljeni skripti3.0, oktober2026. Razlage, vaje z ločenimi rešitvami in opombe predavatelja.'
 slides=[];targets={};idmap={};ranges={}
 for i,s in enumerate(specs):
  slide=prs.slides.add_slide(prs.slide_layouts[6]);slide._idx=i+1;slides.append(slide);idmap[s['id']]=slide
  if s['kind']=='chapter':targets[s['chapter']]=slide
  if s['kind']=='toc':targets['toc']=slide
  if s['id']=='appendix-start':targets['appendix']=slide
  if s.get('chapter'):
   n=s['chapter'];ranges.setdefault(n,[i+1,i+1]);ranges[n][1]=i+1
 covered=set();notes_words=0
 for i,(slide,s) in enumerate(zip(slides,specs),1):
  frame(slide,s,i,len(slides),targets)
  kind=s['kind']
  if kind in ['cover','chapter','appendix_open']:opening(slide,s,i,len(slides))
  elif kind in ['concept','guide','reading']:concept(slide,s)
  elif kind=='contrast':contrast(slide,s)
  elif kind=='example':example(slide,s)
  elif kind=='process':process(slide,s)
  elif kind=='table':table(slide,s)
  elif kind in ['exercise','answer']:practice(slide,s)
  elif kind=='recap':recap(slide,s)
  elif kind=='toc':toc(slide,s,targets,ranges)
  elif kind=='glossary':glossary(slide,s)
  elif kind=='bibliography':bibliography(slide,s)
  else:raise ValueError(kind)
  notes_words+=make_notes(slide,s,source,covered)
 add_sections(prs,specs)
 normalize_app_metadata(prs,specs)
 out=args.output or ROOT/('LSSJ_sample.pptx' if args.sample else 'LSSJ_predavanja.pptx')
 out.parent.mkdir(parents=True,exist_ok=True)
 prs.save(out)
 expected_sections={k for n,c in source.items() for k in c['sections'] if k not in [c['exercise_section'],c['solution_section']]}
 expected_tasks={(int(n),x['number']) for n,c in source.items() for x in c['exercises']}
 projected_tasks={(s['chapter'],x['number']) for s in specs if s['kind']=='exercise' for x in s['items']}
 projected_answers={(s['chapter'],x['number']) for s in specs if s['kind']=='answer' for x in s['items']}
 if not args.sample:
  assert not(expected_sections-covered),('Missing source sections',expected_sections-covered)
  assert projected_tasks==projected_answers==expected_tasks
  assert len({s['id'] for s in specs})==len(specs)
 report={'file':str(out),'slides':len(slides),'bytes':out.stat().st_size,'chapter_ranges':ranges,'exercise_sets':len(projected_tasks),'answer_sets':len(projected_answers),'hard_sets':sum(len(s['items']) for s in specs if s['kind']=='exercise' and s['hard']),'glossary_entries':termcount,'notes_words':notes_words,'source_sections_covered':sorted(covered),'kinds':{k:sum(s['kind']==k for s in specs) for k in sorted({s['kind'] for s in specs})},'layout_warnings':WARN}
 report['measurement_font']=Path(FONT_FILE).name
 report['presentation_font']=FONT
 if args.report:
  args.report.parent.mkdir(parents=True,exist_ok=True)
  args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
